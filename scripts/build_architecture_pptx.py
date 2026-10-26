#!/usr/bin/env python3
"""Regenerate the editable Mind Blogging architecture reference deck.

Source-of-truth: docs/architecture-reference.md and repository code.
Usage: pip install python-pptx==1.0.2 && python scripts/build_architecture_pptx.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "Mind_Blogging_Tech_Stack_Architecture.pptx"
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
W, H = 13.333, 7.5
BG = (11, 17, 35)
CARD = (24, 35, 59)
CYAN = (65, 216, 223)
LILAC = (169, 145, 255)
MUTED = (169, 185, 210)
WHITE = (248, 250, 253)
ORANGE = (255, 194, 116)
GREEN = (134, 224, 174)

def rgb(t):
    return RGBColor(*t)

def shape(slide, x, y, w, h, fill=CARD, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.fill.background() if line is None else None
    if line is not None:
        item.line.color.rgb = rgb(line)
    return item

def txt(slide, content, x, y, w, h, size=20, color=WHITE, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = box.text_frame
    f.clear()
    f.word_wrap = True
    f.margin_left = f.margin_right = Inches(0.02)
    f.margin_top = f.margin_bottom = Inches(0.02)
    for j, line in enumerate(str(content).split("\\n")):
        para = f.paragraphs[0] if j == 0 else f.add_paragraph()
        para.text = line
        para.space_after = Pt(7 if size <= 18 else 4)
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box

def slide(title, section, description="", number=None):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    shape(s, 0, 0, W, 0.10, CYAN, False)
    txt(s, section.upper(), 0.62, 0.40, 11.7, 0.35, 12, CYAN, True)
    txt(s, title, 0.62, 0.95, 12.1, 0.66, 30, WHITE, True)
    if description:
        txt(s, description, 0.64, 1.72, 11.9, 0.55, 15, MUTED)
    txt(s, "MIND BLOGGING  /  TECHNICAL REFERENCE", 0.63, 7.12, 9.8, 0.25, 10, MUTED)
    txt(s, str(number or len(P.slides)).zfill(2) + " / 18", 11.70, 7.11, 0.9, 0.25, 10, MUTED)
    return s

def panel(s, title, body, x, y, w, h, accent=CYAN, title_size=18, body_size=14):
    shape(s, x, y, w, h)
    shape(s, x, y, 0.065, h, accent, False)
    txt(s, title, x+0.26, y+0.22, w-0.51, 0.48, title_size, accent, True)
    txt(s, body, x+0.26, y+0.84, w-0.55, h-0.96, body_size, WHITE)
    
def list_slide(title, section, description, cards, number=None):
    s = slide(title, section, description, number)
    n = len(cards)
    if n == 2:
        y, h, gap = 2.50, 3.60, 0.28
        width = (12.05-gap)/2
        for i,(head,body) in enumerate(cards):
            panel(s,head,body,0.65+i*(width+gap),y,width,h,(CYAN,LILAC)[i])
    elif n == 3:
        y, h, gap = 2.50, 3.60, 0.20
        width = (12.05-2*gap)/3
        for i,(head,body) in enumerate(cards):
            panel(s,head,body,0.65+i*(width+gap),y,width,h,(CYAN,LILAC,ORANGE)[i])
    elif n == 4:
        w, h = 5.91, 1.70
        for i,(head,body) in enumerate(cards):
            panel(s,head,body,0.65+(i%2)*6.15,2.38+(i//2)*1.97,w,h,(CYAN,LILAC,ORANGE,GREEN)[i],16,13)
    else:
        raise ValueError("Choose two, three or four cards")
    return s

s=slide("Mind Blogging", "Project deep dive  •  October 2026", "Tech Stack & How It Works  |  A code-grounded, editable reference", 1)
txt(s,"REACT  /  UI",1.05,3.0,3.4,0.5,24,CYAN,True)
txt(s,"SUPABASE  /  DATA",4.80,3.0,3.8,0.5,24,LILAC,True)
txt(s,"VERCEL  /  SERVICES",8.95,3.0,3.7,0.5,24,ORANGE,True)
shape(s,1.06,4.04,11.15,0.02,MUTED,False)
txt(s,"From components and sign-in to persistence, AI streaming, reminders, security and deployment.",1.04,4.5,11.2,1.2,22,WHITE)

list_slide("Two paths, one app", "01  /  Mental model",
           "A key distinction: the journal adapter does not use the Vercel API.",
           [("PATH A  •  RECORDS","React → src/cloud.ts → Supabase\\nCreate, edit, read and delete entries, goals, reviews and settings."),
            ("PATH B  •  SERVER TASKS","React → /api/*.js → Supabase + external services\\nAI inference, suggested tags, proposal actions and push delivery.")],2)

list_slide("What is the app built with?", "02  /  Technology stack",
           "Major dependencies and external services verified from the repository.",
           [("FRONTEND","React 19  •  TypeScript 5.7\\nVite 6  •  Tailwind CSS 3.4\\nCustom CSS  •  Lucide React"),
            ("DATA & SECURITY","Supabase PostgreSQL\\nSupabase Auth  •  User sessions\\nRow Level Security is required."),
            ("SERVICES & DELIVERY","Vercel serverless functions\\nVercel AI Gateway  •  Vercel Blob\\nWeb Push  •  Vercel Cron  •  GitHub")],3)

list_slide("Under the hood: request routing", "03  /  System architecture",
           "Keep the client, data plane and server-only credentials separate.",
           [("BROWSER LAYER","React components + hooks\\nAuth sessions in Supabase client\\nLocal record adapter: src/cloud.ts"),
            ("SUPABASE LAYER","Auth: validates user identity\\nPostgreSQL: mind_blog_records\\nRLS: isolate users by user_id"),
            ("VERCEL LAYER","/api/ai: AI Gateway streaming\\n/api/push: private subscription storage\\n/api/cron: authenticated schedule")],4)

list_slide("Where to find things", "04  /  Codebase map",
           "Start with these files when debugging a feature.",
           [("src/  •  BROWSER","main.tsx  •  App.tsx\\nPlayfulProfile.tsx  •  PlayfulLogin.tsx\\nPlayfulWorkspace.tsx  •  PlayfulAdvanced.tsx"),
            ("api/  •  SERVER","_shared.js: token verification\\nai.js: model gateway\\naction.js  •  suggest-tags.js"),
            ("OTHER IMPORTANT FILES","cloud.ts: data and auth adapter\\npush.js  •  cron.js  •  public/sw.js\\nvercel.json  •  tests/release.test.mjs")],5)

list_slide("How React starts and renders", "05  /  Frontend lifecycle",
           "The currently mounted app uses the Playful UI components.",
           [("1  •  BOOTSTRAP","index.html exposes #root.\\nsrc/main.tsx mounts React <App/>."),
            ("2  •  APP ENTRY","src/App.tsx first renders PlayfulProfile.\\nEntering a page mounts PlayfulWorkspace."),
            ("3  •  UI STATE","useState manages draft and page state.\\nuseEffect loads data and listens for auth changes.")],6)

list_slide("Authentication: click to session", "06  /  Identity",
           "The UI offers multiple methods; Google needs provider configuration.",
           [("START","User selects password, email link or Google sign-in in PlayfulLogin."),
            ("VERIFY","src/cloud.ts calls Supabase Auth. A successful sign-in creates a user session."),
            ("AUTHORIZE","Client forwards access token for serverless API calls; backend revalidates the user.")],7)

list_slide("The database + data adapter", "07  /  Persistence",
           "Most CRUD goes directly to Supabase, without a Vercel serverless hop.",
           [("TABLE  •  mind_blog_records","Composite identity: user_id + kind + id\\nPayload: record body / JSON data\\nKinds: entries, artifacts, goals, reviews, settings"),
            ("ADAPTER  •  src/cloud.ts","read(): filters records by signed-in user\\nwrite(): upsert record by composite key\\nremove(): delete matching owned record")],8)

list_slide("Follow a journal entry save", "08  /  Data flow",
           "What happens after somebody clicks Save?",
           [("IN THE BROWSER","1  User edits a journal form\\n2  React holds the current draft\\n3  PlayfulWorkspace calls api.post('/api/entries', …)"),
            ("IN THE DATA ADAPTER","4  cloud.ts handles the call locally\\n5  Supabase upsert writes the record\\n6  React reloads / updates its UI")],9)

list_slide("How a reflection becomes an answer", "09  /  AI execution",
           "Consent and data ownership are enforced before the gateway request.",
           [("REQUEST + CHECK","POST /api/ai with bearer token\\nValidate user and requested entry\\nCheck aiConsent and aiExcluded"),
            ("GATEWAY + RESPONSE","AI_GATEWAY_API_KEY stays on server\\nGateway streams completion to /api/ai\\nServer writes returned text to the UI")],10)

list_slide("How browser notifications are delivered", "10  /  Reminders",
           "Web Push links the browser, private subscriptions and an authenticated cron.",
           [("OPT-IN","Browser grants notifications and registers sw.js.\\nSubscription saved to private Vercel Blob."),
            ("SCHEDULE","vercel.json runs Monday 03:30 UTC / 09:00 IST.\\n/api/cron requires CRON_SECRET."),
            ("DELIVERY","Backend sends Web Push messages.\\nService worker shows notifications and handles clicks.")],11)

list_slide("Safeguards in the code", "11  /  Security",
           "Protection must work at multiple layers, not just in visible UI.",
           [("IDENTITY","Supabase authentication\\nBearer-token validation\\nRecord ownership checks"),
            ("DATA PRIVACY","RLS needed on mind_blog_records\\nAI consent and per-entry exclusion\\nPrivate keys kept out of browser"),
            ("OPERATIONS","Cron fails closed without its secret\\nVersion checks on edits\\nAccount erasure cleans up push records")],12)

list_slide("How code reaches users", "12  /  Shipping",
           "The frontend and backend deploy to Vercel; Supabase is a separate service.",
           [("SOURCE","Developer changes code\\nCommits and pushes to GitHub"),
            ("BUILD","Vercel builds the Vite frontend\\n/api files become serverless routes"),
            ("LIVE APP","Visitors load compiled React assets\\nProduction data/auth remain in Supabase")],13)

list_slide("Run, inspect and test locally", "13  /  Developer workflow",
           "Node.js 20+ is required by project documentation.",
           [("SETUP","git clone https://github.com/jd2009-k/mind-blogging.git\\ncd mind-blogging\\nnpm install"),
            ("COMMANDS","npm run dev\\nnpm run check\\nnpm run test\\nnpm run build"),
            ("IMPORTANT LIMITATION","Local Vite serves the frontend.\\nUse Vercel-compatible execution for /api serverless functions.")],14)

list_slide("What is tested — and what is not", "14  /  Verification",
           "Source: QA_REPORT_2026-10-26.md, reflecting the preview deployment.",
           [("REPORTED PASS","TypeScript + Vite build checks\\nSource-regression tests\\nBasic guest HTTP / API security smoke tests"),
            ("NOT END-TO-END VERIFIED","Actual registration/sign-in and database CRUD\\nSigned-in AI streaming / cancellation\\nDevice push delivery"),
            ("RELEASE CONTEXT","Preview was READY according to report.\\nThis tested preview was not promoted to production.")],15)

list_slide("What still needs attention", "15  /  Release priorities",
           "Source checks alone cannot establish complete production readiness.",
           [("P0  •  AUTH + DATA","Confirm signup, email delivery and session persistence.\\nTest CRUD and RLS with two test accounts."),
            ("P1  •  AI + NOTIFICATIONS","Confirm authenticated AI streaming and cancel behavior.\\nVerify service-worker push receipt on devices."),
            ("P2  •  DELIVERY","Commit lockfile and verify npm ci.\\nAutomate dependency audit and regression QA.")],16)

list_slide("How to learn the codebase", "16  /  Learning path",
           "Move from programming basics to actual request tracing.",
           [("PHASE 1  •  LANGUAGE","JavaScript fundamentals and async/await\\nTypeScript types and interfaces\\nReact components, hooks and state"),
            ("PHASE 2  •  APPLICATION","Read App.tsx and PlayfulWorkspace.tsx\\nTrace cloud.ts and Supabase records\\nUnderstand auth and RLS"),
            ("PHASE 3  •  OPERATIONS","Study api/ai.js, push.js and cron.js\\nRun tests and deployment checks\\nMake a small feature change safely")],17)

list_slide("Know the flow. Own the code.", "The big picture",
           "One app, three critical responsibilities.",
           [("REACT","Renders the interface and manages interaction state."),
            ("SUPABASE","Authenticates users and stores private user records."),
            ("VERCEL","Hosts the frontend and executes privileged AI / push tasks.")],18)

OUT.parent.mkdir(parents=True, exist_ok=True)
P.save(OUT)
print(f"Created {OUT} ({len(P.slides)} slides)")
